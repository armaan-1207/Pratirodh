"""Create the main Word guide from its Markdown source using python-docx."""
import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def populate_results():
    path = ROOT / 'docs/PROJECT_GUIDE.md'
    text = path.read_text(encoding='utf-8').split('## Measured release results')[0]
    result = json.loads((ROOT / 'run_output/benchmark-upgrade.json').read_text())
    lines = ['## Measured release results', '', f"The expanded comparison recorded {len(result['rows'])} evaluations on a catalogue of 144 curated candidates across 36 synthetic scenarios and four modes. These are related constructed families, not production accuracy. Cloud generation was not invoked.", '']
    for mode, values in result['metrics'].items():
        lines.append(f"- {mode.capitalize()} verification: {values['correct_ready']} of {values['correct_total']} correct patches ready; {values['incorrect_ready']} of {values['incorrect_total']} incorrect patches incorrectly ready; {values['abstentions']} abstentions; {values['seconds']} cumulative seconds.")
    local_path = ROOT / 'run_output/local-evaluation.json'
    if local_path.exists():
        local = json.loads(local_path.read_text())
        lines += ['', 'Real local Ollama generation is a separate cohort. Every attempted scenario is listed below; unsuccessful outcomes are retained.']
        for row in local['rows']:
            lines.append(f"- {row['scenario']}: {row['decision']}; {row['seconds']} seconds; evidence {row['run_id']}.")
    lines += ['', 'Full and unguided modes each spend 64 additional request executions per candidate; qualification counts against the guided budget. This implementation uses a shared trusted input pool, so similar outcomes are possible. No superiority claim follows without measured differences.', '', 'The original 0.1 comparison contained 288 evaluations: full mode accepted 24 correct patches and no incorrect patch out of 72. Those historical results are kept separately and must not be substituted for this release.', '', 'The final check record is docs/RELEASE_CHECKS.json. The operator applied program-specific Windows firewall rules. The dedicated Python executable retained loopback access and its external connection was denied; a control interpreter retained external access. Ollama has an active outbound rule and disabled cloud features. The real cohort uses the dedicated interpreter and an additional Python audit policy.', '', 'A narrated MP4 and organizer-template PPT remain team deliverables; this guide supplies their content and recording sequence.', '']
    path.write_text(text + '\n'.join(lines), encoding='utf-8')


def add_runs(paragraph, text):
    for i, part in enumerate(re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)):
        run = paragraph.add_run(part.strip('*`') if part.startswith(('**', '`')) else part)
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Liberation Mono'
            run.font.size = Pt(9)


def main():
    populate_results()
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = document.styles[style_name]
        style.font.name = 'Liberation Sans'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + key), None)
        fonts.set(qn('w:ascii'), 'Liberation Sans')
        fonts.set(qn('w:hAnsi'), 'Liberation Sans')
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    document.styles['Title'].font.size = Pt(26)
    document.styles['Heading 1'].font.size = Pt(16)
    document.styles['Heading 2'].font.size = Pt(12.5)
    for name in ['Heading 1', 'Heading 2']:
        document.styles[name].paragraph_format.keep_with_next = True
        document.styles[name].paragraph_format.space_before = Pt(14)
        document.styles[name].paragraph_format.space_after = Pt(6)
    code_style = document.styles.add_style('Code', 1)
    code_style.font.name = 'Liberation Mono'
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_after = Pt(2)
    code_style.paragraph_format.line_spacing = 1
    in_code = False
    buffer = []
    def flush():
        if buffer:
            paragraph = document.add_paragraph()
            add_runs(paragraph, ' '.join(buffer))
            buffer.clear()
    for line in (ROOT / 'docs/PROJECT_GUIDE.md').read_text(encoding='utf-8').splitlines():
        if line.startswith('```'):
            flush()
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph(line, 'Code')
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F2F3F4')
            paragraph._p.get_or_add_pPr().append(shade)
            continue
        if not line.strip():
            flush()
        elif re.match(r'^!\[.*\]\(.+\)$', line):
            flush()
            match = re.match(r'^!\[(.*)\]\((.+)\)$', line)
            image_path = (ROOT / 'docs' / match.group(2)).resolve()
            if not image_path.is_relative_to(ROOT / 'docs'):
                raise ValueError('guide image escapes documentation directory')
            document.add_picture(str(image_path), width=Inches(6.5))
            caption = document.add_paragraph(match.group(1))
            caption.paragraph_format.space_after = Pt(12)
            for run in caption.runs:
                run.font.size = Pt(9)
        elif line.startswith('# '):
            flush()
            document.add_paragraph(line[2:], 'Title')
            document.add_paragraph('Technical documentation and demonstration planning', 'Subtitle')
            document.styles['Subtitle'].font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith('## '):
            flush()
            document.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            flush()
            document.add_heading(line[4:], level=2)
        elif line.startswith('- '):
            flush()
            add_runs(document.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            flush()
            add_runs(document.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        else:
            buffer.append(line)
    flush()
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run('PRATIRODH Project Guide  |  ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.title = 'PRATIRODH Project Guide'
    document.core_properties.subject = 'Implementation deployment security demo video and presentation'
    document.core_properties.author = 'PRATIRODH team'
    document.save(ROOT / 'docs/PROJECT_GUIDE.docx')
    print('Created detailed Markdown and Word guides')


if __name__ == '__main__':
    main()
